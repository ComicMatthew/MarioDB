import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl

from app.config import load_config


def _get_temp_dir() -> Path:
    root_dir = Path(__file__).resolve().parent.parent
    temp_dir = root_dir / "temp"
    temp_dir.mkdir(parents=True, exist_ok=True)
    return temp_dir


def _copy_to_temp_file(source_path: str | Path) -> Path:
    source = Path(source_path)
    temp_dir = _get_temp_dir()
    temp_file = temp_dir / f"{source.stem}_{datetime.now().strftime('%Y%m%d%H%M%S%f')}{source.suffix}"
    shutil.copy2(source, temp_file)
    return temp_file


def _replace_file_with_temp_copy(temp_file: str | Path, target_path: str | Path) -> None:
    temp_path = Path(temp_file)
    target = Path(target_path)
    if temp_path.exists():
        shutil.copy2(temp_path, target)


def _format_value(value):
    if value is None:
        return "brak"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def get_used_file(usage_path):
    folder = Path(usage_path)
    files = [f for f in folder.iterdir() if f.suffix.lower() in {".xlsx", ".xlsm"}]
    if not files:
        raise FileNotFoundError(f"W folderze {folder} nie ma zadnego pliku do przetworzenia!")

    todo_file = files[0]
    return str(todo_file), todo_file.name


def _read_config_values():
    config = load_config()
    return {
        "database_file_path": config.get("database_file_path", ""),
        "usage_subtraction_file_path": config.get("usage_subtraction_file_path", ""),
        "usage_addition_file_path": config.get("usage_addition_file_path", ""),
        "done_folder_path": config.get("done_folder_path", ""),
        "usage_worksheet": config.get("usage_worksheet", ""),
        "usage_workdatasheet": config.get("usage_workdatasheet", ""),
        "substraction_start_row": int(config.get("substraction_start_row", 0) or 0),
        "addition_start_row": int(config.get("addition_start_row", 0) or 0),
        "database_start_row": int(config.get("database_start_row", 0) or 0),
    }


def find_missing_records(database_path, usage_path, usage_worksheet, database_worksheet, start_row, database_start_row):
    temp_database_path = _copy_to_temp_file(database_path)
    database_wb = openpyxl.load_workbook(temp_database_path)
    try:
        todo_file_path, todo_file_name = get_used_file(usage_path)
        usage_wb = openpyxl.load_workbook(todo_file_path)
        database_sheet = database_wb[database_worksheet]
        usage_sheet = usage_wb[usage_worksheet]
        missing = []

        for usage_row in usage_sheet.iter_rows(min_row=start_row, values_only=True):
            if not usage_row:
                continue
            asset_name = usage_row[1] if len(usage_row) > 1 else None
            if asset_name is None:
                continue

            found = False
            for database_row in database_sheet.iter_rows(min_row=database_start_row, max_row=database_sheet.max_row, values_only=True):
                if not database_row:
                    continue
                if database_row[1] == asset_name:
                    found = True
                    break

            if not found:
                missing.append(asset_name)

        return missing, todo_file_name
    finally:
        database_wb.close()
        if temp_database_path.exists():
            temp_database_path.unlink()


def update_quantities(database_path, usage_path, done_folder_path, usage_worksheet, database_worksheet, start_row, database_start_row, operation="add"):
    temp_database_path = _copy_to_temp_file(database_path)
    database_wb = openpyxl.load_workbook(temp_database_path)
    try:
        todo_file_path, todo_file_name = get_used_file(usage_path)
        usage_wb = openpyxl.load_workbook(todo_file_path)

        database_sheet = database_wb[database_worksheet]
        usage_sheet = usage_wb[usage_worksheet]
        count = 0
        negative_record = []
        audit_entries = []
        operation_label = "dodano" if operation == "add" else "odjęto"

        for row in usage_sheet.iter_rows(min_row=start_row, values_only=True):
            if not row or len(row) <= 1:
                continue
            asset_name, used_quantity = row[1], row[3] if len(row) > 3 else None

            for index, database_row in enumerate(database_sheet.iter_rows(min_row=database_start_row, max_row=database_sheet.max_row, values_only=True), start=2):
                if not database_row:
                    continue
                asset_number, current_quantity = database_row[1], database_row[3] if len(database_row) > 3 else None
                if asset_number == asset_name and current_quantity is not None and used_quantity is not None:
                    if operation == "subtract":
                        new_quantity = current_quantity - used_quantity
                    else:
                        new_quantity = current_quantity + used_quantity

                    database_sheet.cell(row=index, column=4).value = new_quantity
                    count += 1
                    audit_entries.append({
                        "record": asset_name,
                        "from": _format_value(current_quantity),
                        "to": _format_value(new_quantity),
                        "source_file": todo_file_name,
                        "message": (
                            f"Rekord {asset_name} zmienił wartość z {_format_value(current_quantity)} "
                            f"na {_format_value(new_quantity)} ({operation_label} {used_quantity} na podstawie pliku {todo_file_name})."
                        ),
                    })
                    if new_quantity < 0:
                        negative_record.append(asset_number)

        database_wb.save(temp_database_path)
        _replace_file_with_temp_copy(temp_database_path, database_path)
        new_name = todo_file_name.split(".")[0] + datetime.now().strftime(" %Y-%m-%d %H-%M-%S") + "." + todo_file_name.split(".")[-1]
        done_path = Path(done_folder_path)
        done_path.mkdir(parents=True, exist_ok=True)
        final_done_file = done_path / new_name.strip()
        shutil.move(todo_file_path, str(final_done_file))

        summary = (
            f"Zaktualizowano {count} rekordów w trybie {operation_label}. "
            f"Plik {todo_file_name} został przetworzony i przeniesiony do {final_done_file.name}."
        )

        return {
            "updated": count,
            "negative_records": negative_record,
            "done_file": str(final_done_file),
            "audit": audit_entries,
            "summary": summary,
        }
    finally:
        database_wb.close()
        if temp_database_path.exists():
            temp_database_path.unlink()


def process_subtraction():
    config = _read_config_values()
    missing, _ = find_missing_records(
        config["database_file_path"],
        config["usage_subtraction_file_path"],
        config["usage_worksheet"],
        config["usage_workdatasheet"],
        config["substraction_start_row"],
        config["database_start_row"],
    )
    result = update_quantities(
        config["database_file_path"],
        config["usage_subtraction_file_path"],
        config["done_folder_path"],
        config["usage_worksheet"],
        config["usage_workdatasheet"],
        config["substraction_start_row"],
        config["database_start_row"],
        operation="subtract",
    )
    return {"missing": missing, **result}


def process_addition():
    config = _read_config_values()
    missing, _ = find_missing_records(
        config["database_file_path"],
        config["usage_addition_file_path"],
        config["usage_worksheet"],
        config["usage_workdatasheet"],
        config["addition_start_row"],
        config["database_start_row"],
    )
    result = update_quantities(
        config["database_file_path"],
        config["usage_addition_file_path"],
        config["done_folder_path"],
        config["usage_worksheet"],
        config["usage_workdatasheet"],
        config["addition_start_row"],
        config["database_start_row"],
        operation="add",
    )
    return {"missing": missing, **result}
